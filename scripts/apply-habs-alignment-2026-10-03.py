"""
The HABS score, aligned with the TPF app — 2026-10-03.

The owner, 2026-10-03: "make HABS score align". The app's HABS scores nine
components (lower-body strength, power, upper-body push, upper-body pull,
running intensity, running distance, swimming, cycling, rowing / erg) with
its own literal per-pathway weights; this site's composite scored a different
set (docs/HABS-ALIGNMENT-2026-10-03.md, sections 1-3, is the criteria list).
This script is the record of every workbook cell that change needs. It edits,
in place, only:

  config/standards/TPF_HRS_Standards_v0_2026-06-21.xlsx  (Lift / Hybrid)

Then run `npm run codegen` (the generated files are never hand-edited) and
check against the app:

  npm run check:app-habs -- <path to tpf-app>   (the HABS model: weights,
                                                 membership, ladders, scores)
  npm run check:app-lift -- <path to tpf-app>   (the ladders, incl. the new
                                                 derived 10 km / half)

WHAT IT CHANGES

  1. Benchmarks_Sourcing gains a `habs_component` column: which of the app's
     nine HABS components a benchmark feeds (tpf-app src/lib/habs.ts
     COMPONENT_DEFS). Blank = the benchmark is a TPF Benchmark standard that is
     NOT in the HABS score (front squat, snatch, clean & jerk, the 500 m row,
     broad jump, the gymnastics rows, plank, and the optional grip / ruck).
     The existing `component` column is untouched: it still groups the
     Capacity Index (Quality_Mix) and the Operator-style components.
  2. Two benchmarks the app scores and this site never had: the 10 km run and
     the half marathon (`run_10k`, `run_half`, component `running`, HABS
     component `run_distance`). Their ladders are DERIVED, exactly as the app
     derives them at runtime (habs.ts riegelStd): the 5 km ladder x
     (d / 5000) ^ 1.06, to the nearest 10 s, a half rounding up. The base
     rows come from the base 5 km; gym_goer, crossfit_generalist and
     triathlete — the three pathways with their own 5 km — get their own
     derived rows on Standards_Pathway (habs_pathway_standards.ts
     addDerivedRunKeys). The vestigial four-tier `excellent` is derived the
     same way from the 5 km's excellent (it is never scored on a six-tier row).
  3. A new HABS_Weights sheet: the app's literal HABS_PATHWAY_WEIGHTS
     (tpf-app src/lib/habs_pathways.ts), component x pathway, each column
     summing to 100. These are the app's numbers, not a re-derivation: the app
     made them from this workbook's Weights sheet (dropping olympic,
     gymnastics and core endurance, splitting upper 13 : 12 and running
     50 : 50, renormalising) and then hand-rounded two of them (crossfit power
     19.2, bodybuilder push 22.7) so each row totals 100. The existing Weights
     sheet is kept: it still decides which extra standards a pathway lists and
     feeds the Capacity Index.
  4. A Read me row recording all of the above.

HOW IT IS BUILT (the guards of scripts/apply-standards-research-2026-10-03.py)

  - A derived row or weight cell that already exists must hold exactly the
    value this script would write; anything else stops the script before
    anything is saved (the 5 km moved, or someone edited a derived row —
    a person should look).
  - A `habs_component` cell that holds a different non-blank value stops it.
  - A `verify` pass runs on the finished sheets before saving.
  - Idempotent: a re-run finds everything in place and saves nothing.

Requires: python3 with openpyxl (3.1.5 checked).
"""

import math
from copy import copy
from pathlib import Path

import openpyxl

REPO = Path(__file__).resolve().parent.parent
HRS = REPO / 'config/standards/TPF_HRS_Standards_v0_2026-06-21.xlsx'

changes = []
dirty = False


def put(ws, row, col, value):
    global dirty
    cell = ws.cell(row=row, column=col)
    if cell.value != value:
        changes.append(f'{ws.title}!{cell.coordinate}: {cell.value!r} -> {value!r}')
        cell.value = value
        dirty = True


def header_index(ws, header_row):
    return {str(c.value).strip().lower(): c.column for c in ws[header_row] if c.value is not None}


def round_half_up(x, step):
    """Nearest `step`, a half rounding UP — JS Math.round, which the app uses
    (Python's round() rounds a half to even)."""
    return int(math.floor(x / step + 0.5)) * step


def to_sec(v):
    if isinstance(v, (int, float)):
        return v
    out = 0
    for part in str(v).split(':'):
        out = out * 60 + int(part)
    return out


def mmss(sec):
    """Seconds -> the workbook's m:ss form (minutes may exceed 59, as on the
    40 km bike rows)."""
    sec = int(sec)
    return f'{sec // 60}:{sec % 60:02d}'


def col_of(idx, prefix):
    for name, c in idx.items():
        if name.startswith(prefix):
            return c
    raise SystemExit(f'no column starting {prefix!r}')


wb = openpyxl.load_workbook(HRS)
errors = []

# ------------------------------------------------------------- the app ------
# tpf-app src/lib/habs.ts — COMPONENT_DEFS order, and the Riegel factors.
HABS_COMPONENTS = ['lower_strength', 'power', 'upper_push', 'upper_pull',
                   'run_intensity', 'run_distance', 'swimming', 'cycling', 'erg']
RIEGEL_10K = (10000 / 5000) ** 1.06       # 2.084931521682243 — the app's, to the bit
RIEGEL_HALF = (21097.5 / 5000) ** 1.06    # 4.600199333256586
PATHWAYS = ['gym_goer', 'hybrid_athlete', 'crossfit_generalist', 'hyrox', 'powerlifter', 'bodybuilder', 'triathlete']

# tpf-app src/lib/habs_pathways.ts HABS_PATHWAY_WEIGHTS, verbatim (2026-10-03, 226b018a).
HABS_WEIGHTS = {
    'gym_goer':            dict(lower_strength=28.6, power=14.3, upper_push=14.9, upper_pull=13.7,
                                run_intensity=7.1, run_distance=7.1, swimming=0, cycling=0, erg=14.3),
    'hybrid_athlete':      dict(lower_strength=20, power=13.4, upper_push=10.4, upper_pull=9.6,
                                run_intensity=13.3, run_distance=13.3, swimming=0, cycling=0, erg=20),
    'crossfit_generalist': dict(lower_strength=22.6, power=19.2, upper_push=10.1, upper_pull=9.3,
                                run_intensity=9.7, run_distance=9.7, swimming=0, cycling=0, erg=19.4),
    'hyrox':               dict(lower_strength=20.8, power=9.1, upper_push=5.4, upper_pull=5.0,
                                run_intensity=21.4, run_distance=21.4, swimming=0, cycling=0, erg=16.9),
    'powerlifter':         dict(lower_strength=50, power=11.1, upper_push=20.2, upper_pull=18.7,
                                run_intensity=0, run_distance=0, swimming=0, cycling=0, erg=0),
    'bodybuilder':         dict(lower_strength=43.8, power=12.5, upper_push=22.7, upper_pull=21.0,
                                run_intensity=0, run_distance=0, swimming=0, cycling=0, erg=0),
    'triathlete':          dict(lower_strength=10, power=5, upper_push=5, upper_pull=5,
                                run_intensity=12.5, run_distance=12.5, swimming=25, cycling=25, erg=0),
}

# Site benchmark id -> the app's HABS component (habs.ts COMPONENT_DEFS).
HABS_OF = {
    'back_squat_1rm': 'lower_strength', 'deadlift_1rm': 'lower_strength',
    'power_clean_1rm': 'power',
    'bench_1rm': 'upper_push', 'strict_press_1rm': 'upper_push',
    'barbell_row_1rm': 'upper_pull',
    'run_1mi': 'run_intensity', 'run_5k': 'run_intensity',
    'run_10k': 'run_distance', 'run_half': 'run_distance',
    'swim_400m': 'swimming', 'swim_1500m': 'swimming',
    'bike_20k': 'cycling', 'bike_40k': 'cycling',
    'row_2k': 'erg',
}

TIERS = ['pass', 'novice', 'good', 'excellent', 'intermediate', 'advanced', 'elite']
DERIVED = (('run_10k', RIEGEL_10K, '10 km'), ('run_half', RIEGEL_HALF, 'half marathon'))
DERIVED_SRC = ("Derived: TPF's 5 km ladder x (d / 5000) ^ 1.06 (Riegel), to the nearest 10 s — the TPF app's "
               "runtime derivation (habs.ts riegelStd), mirrored 2026-10-03")
DERIVED_NOTE = 'Derived — change the 5 km, not this (2026-10-03)'

# ------------------------------------------------- 1. Benchmarks_Sourcing ---

src = wb['Benchmarks_Sourcing']
sidx = header_index(src, 2)
if 'habs_component' not in sidx:
    put(src, 2, src.max_column + 1, 'habs_component')
    src.cell(row=2, column=src.max_column)._style = copy(src.cell(row=2, column=src.max_column - 1)._style)
    sidx = header_index(src, 2)


def sourcing_rows():
    return {src.cell(row=r, column=sidx['benchmark_id']).value: r
            for r in range(3, src.max_row + 1) if src.cell(row=r, column=sidx['benchmark_id']).value}


rows = sourcing_rows()
run5 = rows['run_5k']
SOURCING = {
    'run_10k': ('running', 'race_times', 'mm:ss', True, 'absolute',
                "TPF's own: derived from TPF's 5 km ladder (run_5k) with Riegel's formula, exponent 1.06, to the "
                "nearest 10 s — the same derivation the TPF app's HABS uses at runtime (2026-10-03)",
                "TPF's own", 'Yes', 'Trained adults (hybrid)', 'Derived from run_5k (2026-10-03)', DERIVED_NOTE),
    'run_half': ('running', 'race_times', 'mm:ss', True, 'absolute',
                 "TPF's own: derived from TPF's 5 km ladder (run_5k) with Riegel's formula, exponent 1.06, to the "
                 "nearest 10 s — the same derivation the TPF app's HABS uses at runtime (2026-10-03)",
                 "TPF's own", 'Yes', 'Trained adults (hybrid)', 'Derived from run_5k (2026-10-03)', DERIVED_NOTE),
}
SRC_COLS = ['component', 'source', 'unit', 'lower_is_better', 'normalization', 'data_source', 'license',
            'commercial_use', 'reference_population', 'launch_method', 'notes']
for offset, (bid, vals) in enumerate(SOURCING.items(), start=1):
    rows = sourcing_rows()
    if bid in rows:
        r = rows[bid]
        for name, v in zip(SRC_COLS, vals):
            have = src.cell(row=r, column=sidx[name]).value
            if have != v:
                errors.append(f'Benchmarks_Sourcing {bid}.{name}: holds {have!r}, this record says {v!r}')
        continue
    at = rows['run_5k'] + offset
    src.insert_rows(at)
    for c in range(1, src.max_column + 1):
        src.cell(row=at, column=c)._style = copy(src.cell(row=at - 1, column=c)._style)
    put(src, at, sidx['benchmark_id'], bid)
    for name, v in zip(SRC_COLS, vals):
        put(src, at, sidx[name], v)

for bid, r in sourcing_rows().items():
    want = HABS_OF.get(bid)
    have = src.cell(row=r, column=sidx['habs_component']).value
    if have not in (None, '') and have != want:
        errors.append(f'Benchmarks_Sourcing {bid}.habs_component: holds {have!r}, this record says {want!r}')
    elif want is not None:
        put(src, r, sidx['habs_component'], want)

# ------------------------------------------- 2a. Standards (base, derived) ---

std = wb['Standards']
tidx = header_index(std, 2)
TCOL = {t: col_of(tidx, t) for t in TIERS}


def std_rows(ws, idx, with_pathway):
    out = {}
    for r in range(3, ws.max_row + 1):
        bid = ws.cell(row=r, column=idx['benchmark_id']).value
        if not bid:
            continue
        sex = ws.cell(row=r, column=idx['sex']).value
        key = (bid, ws.cell(row=r, column=idx['pathway']).value, sex) if with_pathway else (bid, sex)
        out[key] = r
    return out


def derived_ladder(ws, tcol, r5, factor):
    return {t: round_half_up(to_sec(ws.cell(row=r5, column=tcol[t]).value) * factor, 10) for t in TIERS}


def ensure_derived(ws, idx, tcol, key, r5, factor, at, extra):
    """`extra` = {column name: value} for the non-tier cells."""
    want = derived_ladder(ws, tcol, r5, factor)
    rows_now = std_rows(ws, idx, 'pathway' in idx)
    if key in rows_now:
        r = rows_now[key]
        for t in TIERS:
            have = ws.cell(row=r, column=tcol[t]).value
            if have is None or to_sec(have) != want[t]:
                errors.append(f'{ws.title} {key} {t}: holds {have!r}, the 5 km derivation gives {mmss(want[t])}')
        return
    ws.insert_rows(at)
    for c in range(1, ws.max_column + 1):
        ws.cell(row=at, column=c)._style = copy(ws.cell(row=at - 1, column=c)._style)
    for name, v in extra.items():
        put(ws, at, idx[name], v)
    for t in TIERS:
        put(ws, at, tcol[t], mmss(want[t]))


for bid, factor, _ in DERIVED:
    for sex in ('M', 'F'):
        rows_now = std_rows(std, tidx, False)
        r5 = rows_now[('run_5k', sex)]
        last = max(r for (b, _s), r in rows_now.items() if b in ('run_5k', 'run_10k', 'run_half'))
        ensure_derived(std, tidx, TCOL, (bid, sex), r5, factor, last + 1, {
            'benchmark_id': bid, 'sex': sex, 'unit': 'mm:ss', 'lower_is_better': 1,
            'source_ref': DERIVED_SRC, 'notes': DERIVED_NOTE,
        })

# ------------------------------- 2b. Standards_Pathway (own-5 km pathways) ---

sp = wb['Standards_Pathway']
pidx = header_index(sp, 2)
PCOL = {t: col_of(pidx, t) for t in TIERS}
for pathway in ('gym_goer', 'crossfit_generalist', 'triathlete'):
    for bid, factor, _ in DERIVED:
        for sex in ('M', 'F'):
            rows_now = std_rows(sp, pidx, True)
            r5 = rows_now.get(('run_5k', pathway, sex))
            if r5 is None:
                errors.append(f'Standards_Pathway: {pathway} has no own run_5k/{sex} row — the app derives from it')
                continue
            last = max(r for (b, p, _s), r in rows_now.items()
                       if p == pathway and b in ('run_5k', 'run_10k', 'run_half'))
            ensure_derived(sp, pidx, PCOL, (bid, pathway, sex), r5, factor, last + 1, {
                'benchmark_id': bid, 'pathway': pathway, 'sex': sex, 'unit': 'mm:ss', 'lower_is_better': 1,
                'source_ref': DERIVED_SRC + f' — from this pathway\'s own 5 km', 'notes': DERIVED_NOTE,
            })

# --------------------------------------------------------- 3. HABS_Weights ---

TITLE = ("HABS pathway weights — the TPF app's HABS_PATHWAY_WEIGHTS (src/lib/habs_pathways.ts), verbatim · "
         "each column sums to 100 · these, not the Weights sheet, weight the HABS score (2026-10-03)")
NOTE = ("Mirrored from the app, not re-derived: the app made these from the Weights sheet (olympic, gymnastics and "
        "core endurance dropped; upper strength split 13 : 12 into push : pull; running split 50 : 50; renormalised "
        "to 100; then hand-rounded so each column totals 100). Triathlete is owner-set directly. Change them in the "
        "app first, then here, then run npm run check:app-habs. Script: scripts/apply-habs-alignment-2026-10-03.py.")
if 'HABS_Weights' not in wb.sheetnames:
    hw = wb.create_sheet('HABS_Weights', index=wb.sheetnames.index('Weights') + 1)
    weights_ws = wb['Weights']
    put(hw, 1, 1, TITLE)
    hw.cell(row=1, column=1)._style = copy(weights_ws.cell(row=1, column=1)._style)
    for c, name in enumerate(['component'] + PATHWAYS, start=1):
        put(hw, 2, c, name)
        hw.cell(row=2, column=c)._style = copy(weights_ws.cell(row=2, column=min(c, 8))._style)
    for r, comp in enumerate(HABS_COMPONENTS, start=3):
        put(hw, r, 1, comp)
        for c, p in enumerate(PATHWAYS, start=2):
            put(hw, r, c, HABS_WEIGHTS[p][comp])
    total_row = 3 + len(HABS_COMPONENTS)
    put(hw, total_row, 1, 'TOTAL (must = 100)')
    for c, p in enumerate(PATHWAYS, start=2):
        # A number, as on the Weights sheet (a formula has no cached value until Excel recalculates).
        put(hw, total_row, c, round(sum(HABS_WEIGHTS[p].values()), 6))
    put(hw, total_row + 2, 1, NOTE)
    hw.column_dimensions['A'].width = 22
    for c in range(2, 9):
        hw.column_dimensions[openpyxl.utils.get_column_letter(c)].width = 18
else:
    hw = wb['HABS_Weights']
    hidx = header_index(hw, 2)
    found = {}
    for r in range(3, hw.max_row + 1):
        comp = hw.cell(row=r, column=1).value
        if comp in HABS_COMPONENTS:
            found[comp] = r
            for p in PATHWAYS:
                have = hw.cell(row=r, column=hidx[p]).value
                if have != HABS_WEIGHTS[p][comp]:
                    errors.append(f'HABS_Weights {comp}/{p}: holds {have!r}, the app (this record) says '
                                  f'{HABS_WEIGHTS[p][comp]!r}')
    for comp in HABS_COMPONENTS:
        if comp not in found:
            errors.append(f'HABS_Weights: no {comp} row')

# ---------------------------------------------------------------- verify ---

for p in PATHWAYS:
    s = sum(HABS_WEIGHTS[p].values())
    if abs(s - 100) > 0.01:
        errors.append(f'HABS weights {p} sum to {s}, not 100')
for bid in HABS_OF:
    if bid not in sourcing_rows():
        errors.append(f'Benchmarks_Sourcing: no row for {bid}, which feeds HABS {HABS_OF[bid]}')
members = {c: [b for b, h in HABS_OF.items() if h == c] for c in HABS_COMPONENTS}
for c, ids in members.items():
    if not ids:
        errors.append(f'HABS component {c} has no benchmark')
# The derived base rows: the app's numbers (tpf-app habs.ts riegelStd, measured with its tsx 2026-10-03).
APP_BASE = {
    ('run_10k', 'M'): [3760, 3130, 2750, 2440, 2380, 2190], ('run_10k', 'F'): [4340, 3620, 3180, 2830, 2780, 2600],
    ('run_half', 'M'): [8300, 6900, 6070, 5380, 5240, 4830], ('run_half', 'F'): [9570, 7980, 7020, 6230, 6140, 5730],
}
rows_now = std_rows(std, tidx, False)
for key, want in APP_BASE.items():
    r = rows_now.get(key)
    if r is None:
        errors.append(f'Standards: no {key} row')
        continue
    have = [to_sec(std.cell(row=r, column=TCOL[t]).value)
            for t in ('pass', 'novice', 'good', 'intermediate', 'advanced', 'elite')]
    if have != want:
        errors.append(f'Standards {key}: {have}, the app has {want}')

if errors:
    raise SystemExit('A check failed; nothing was saved:\n  ' + '\n  '.join(errors))

# --------------------------------------------------------------- Read me ---

README_KEY = 'Change 2026-10-03 (HABS)'
README = ("The HABS score now matches the TPF app's model (owner: \"make HABS score align\"). New column "
          "Benchmarks_Sourcing.habs_component = which of the app's nine HABS components a benchmark feeds (blank = a TPF "
          "Benchmark standard outside the HABS score). New sheet HABS_Weights = the app's literal per-pathway HABS "
          "weights; the Weights sheet stays for the Capacity Index and for which extra standards a pathway lists. New "
          "benchmarks run_10k and run_half, derived from the 5 km as the app derives them (base, and the three pathways "
          "with their own 5 km). Record: docs/HABS-ALIGNMENT-2026-10-03.md; script: "
          "scripts/apply-habs-alignment-2026-10-03.py.")
rm = wb['Read me']
if not any(c.value == README for c in rm['B']):
    r = rm.max_row + 1
    put(rm, r, 1, README_KEY)
    put(rm, r, 2, README)

# ------------------------------------------------------------------ save ---

if dirty:
    wb.save(HRS)
print(f'{len(changes)} cell change(s); saved: {HRS.name if dirty else "nothing"}')
for c in changes:
    print('  ' + c)
