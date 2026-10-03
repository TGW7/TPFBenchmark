"""
The missing standards, researched — 2026-10-03 (the TPF app's plan 61).

Brings both Excel masters in line with what the TPF app built from its plan 61
on 2026-10-03. The app's record is tpf-app
docs/build/61_MISSING_STANDARDS_RESEARCH_2026-10-03.md section 8 (§8.2 is every
value that moved, old -> new; §8.6 the list of values this repository takes).
The owner's answer, verbatim: "1 yes build all above".

This script is the record of every cell it changes. It edits, in place:

  config/standards/TPF_Operator_Standards.xlsx           (Operator / ORS)
  config/standards/TPF_HRS_Standards_v0_2026-06-21.xlsx  (Lift / Hybrid)

Then run `npm run codegen` — the generated files are never hand-edited — and
check the result against the app:

  npm run check:app-ors  -- <path to tpf-app>   (the Operator units)
  npm run check:app-lift -- <path to tpf-app>   (the WOD ladders, the HABS base
                                                 table and the pathway overrides)

HOW IT IS BUILT (the same guards as scripts/apply-military-top-tiers-2026-10-02.py)

  - Every move names the value it expects to find (OLD) and the value it sets
    (NEW). A cell holding anything else stops the script before ANYTHING is
    saved, in either workbook: the workbook has drifted from what this record
    assumes, and a person should look.
  - A move on the Operator `benchmark` column RENAMES its row. A row is tracked
    by the first name it had, so a re-run (which finds the new name) still
    finds it. A rename never changes the benchmark's id: ids are stored
    (Supabase submissions and logs, the app sync in src/data/appSync.ts), so a
    rename whose `id` cell is blank pins it to the id codegen made from the old
    name.
  - An Operator move's note is appended to the row's `source` column once
    (provenance only; codegen does not emit it). A cite move carries no note:
    the `cite` column is itself the record.
  - Each workbook has a `verify` step run on the finished sheets before
    anything is saved; any error stops the script with nothing saved.
  - Idempotent: a re-run finds every cell at its NEW value, changes nothing and
    saves nothing.

OPERATOR (plan 61 Part B, mirrored for the 13 units this site shares with the app)

  - USAF Pararescue: the app's "IFT (formerly the PAST)" values (USAF IFT
    worksheet, 10 Jan 2023, PJ column — a US government work): 1.5-mile run
    Pass 10:10 -> 10:20; 500 m swim Pass 12:00 -> 12:30 and the row renamed to
    the app's label, which says civilians enlist on the 15:00 column; pull-ups
    renamed "Pull-ups (no time)" -> "Pull-ups (2 min)" (the IFT's own timing;
    1 minute is for officers), its id pinned to `pull_ups_no_time`; push-ups
    (2 min) Pass 50 -> 40; sit-ups (2 min) Pass 54 -> 50. The app's unit source
    label ("USAF Pararescue IFT (formerly the PAST)") has no counterpart here —
    this site has no per-unit source label.
  - Broad jump, one flat ladder per group (plan 61 Part B §3.2): the general
    units (US Navy (PRT), US Police PFT — the app's us_army / air_force /
    uk_army are not mirrored) 1.7 / 1.9 / 2.1 / 2.3 -> 1.70 / 1.95 / 2.20 /
    2.40; the special forces here (SEAL, Pararescue, SFAS — UKSF has no broad
    jump row on this site) 1.8 / 2.0 / 2.2 / 2.4 -> 1.93 / 2.16 / 2.39 / 2.65
    (the AFSPECWAR OFT 8 / 9 / 10-point marks and an Elite about the 80th
    percentile of AFSPECWAR candidates). The combat units and SWAT keep their
    values; their stated basis is mirrored into the `cite` column.
  - Plank: US Police PFT 1:00 / 1:30 / 2:30 / 3:30 -> the operator plank
    1:30 / 2:30 / 3:30 / 4:30 (plan 61 Option A: it sat below every US
    entry-age minimum); US Marine Corps Pass 1:03 -> 1:10 (the PFT minimum
    since 1 Jan 2022). The app's USMC source label now cites MCO 6100.13A (no
    6100.13B was found); this site has no per-unit source label, so that is a
    note only. (The app's firefighter plank moved too; the firefighter unit is
    not mirrored here.)
  - A NEW `cite` column on the Standards sheet holds the app's `cite` verbatim
    for the rows whose stated basis plan 61 wrote — every operator-plank,
    dead-hang and broad-jump row this site holds on a mirrored unit, and the
    Pararescue swim. Codegen does not read it (this site emits no cite);
    `check:app-ors` compares every non-blank cite cell against the app.

LIFT (plan 61 §8.6, and the HABS base table the site's lockstep pins mirror)

  - HYROX race (plan 61 Part A §3.4): men Beginner 1:35:00 -> 1:57:30 and
    Novice 1:30:30 -> 1:37:30 (3.9 x TPF's own 5 km Beginner / Novice); women
    Beginner 1:44:30 -> 2:09:30, Novice 1:39:30 -> 1:47:30, Experienced
    1:35:00 -> 1:34:30 (men x 1.10). Every tier's stated basis rewritten.
  - Cindy: numbers kept, stated basis rewritten (plan 61 Part A §4.4-4.5).
  - Fran, Helen: unchanged.
  - TPF Benchmark's own six-tier broad jump (the app has no such ladder; plan 61
    Part B §3.3): men's Elite 285 -> 275 cm (2.85 sat above every anchor); the
    women's ladder kept (about 0.80 x the men's at every tier); relabelled
    "TPF's own, now checked".
  - THE HABS BASE TABLE (not in plan 61's §8.6 list; found by this run's
    comparison): the app moved three base cells the same day — women's
    overhead press Beginner 20 -> 25 kg (plan 61 Part C §2.2), women's 20 km
    TT Beginner 50:00 -> 49:30 (Part C §4) and the derived 40 km Beginner
    1:43:30 -> 1:42:30. This site's base table is the app's (pathway-
    standards.test.ts "matches the tpf-app base table"; the app is canonical),
    so they move here too. The HYROX, triathlete and bodybuilder pathways' own
    women's overhead press stay at 20 kg, as in the app (its §8.7 Q2).

Requires: python3 with openpyxl (3.1.5 checked).
"""

import math
from pathlib import Path

import openpyxl

REPO = Path(__file__).resolve().parent.parent
OPS = REPO / 'config/standards/TPF_Operator_Standards.xlsx'
HRS = REPO / 'config/standards/TPF_HRS_Standards_v0_2026-06-21.xlsx'

PLAN61 = 'tpf-app mirror 2026-10-03 (plan 61 §8.2)'

changes = []   # printed at the end, each prefixed with its workbook
dirty = set()  # the workbooks (by file name) a change touched; only those are saved


def put(ws, row, col, value):
    cell = ws.cell(row=row, column=col)
    if cell.value != value:
        book = BOOK_NAME[id(ws.parent)]
        changes.append(f'{book} {ws.title}!{cell.coordinate}: {cell.value!r} -> {value!r}')
        dirty.add(book)
        cell.value = value


def header_index(ws, header_row):
    return {str(c.value).strip().lower(): c.column for c in ws[header_row] if c.value is not None}


def round_half_up(x, step):
    """Nearest `step`, a half rounding UP (Python's round() rounds a half to even)."""
    return int(math.floor(x / step + 0.5)) * step


def to_sec(v):
    """A sheet time string (m:ss or h:mm:ss) to seconds; a number passes through."""
    if isinstance(v, (int, float)):
        return v
    out = 0
    for part in str(v).split(':'):
        out = out * 60 + int(part)
    return out


# =========================================================== Operator ======

SEAL, PJ, SFAS = 'Navy SEAL (BUD/S)', 'USAF Pararescue (PJ)', 'US Army Special Forces (SFAS)'
UKSF, PARA, ABN = 'UK Special Forces (SAS/SBS)', 'UK Parachute Regiment (P Coy)', 'US Army Airborne'
RM, USMC, NAVY = 'UK Royal Marines (Cdo Course)', 'US Marine Corps (PFT/CFT)', 'US Navy (PRT)'
POLICE, SWAT = 'US Police PFT', 'US SWAT'
US_INF, UK_INF = 'US Infantry', 'UK Infantry'

IFT_NOTE = ("{what}: the USAF IFT worksheet (10 Jan 2023, PJ column — a US government work), the app's "
            "'USAF Pararescue IFT (formerly the PAST)'; the old values were retired PAST numbers "
            "(plan 61 Part B §2.3) — " + PLAN61)
PJ_SWIM_OLD = '500 m swim (PAST: freestyle, breaststroke or sidestroke)'
PJ_SWIM_NEW = '500 m swim (IFT, PJ column — enlistees 15:00; freestyle, breaststroke or sidestroke)'
PJ_PULL_OLD, PJ_PULL_NEW = 'Pull-ups (no time)', 'Pull-ups (2 min)'

BJ_GENERAL_NOTE = ("Broad jump 1.7 / 1.9 / 2.1 / 2.3 -> 1.70 / 1.95 / 2.20 / 2.40 m: the general flat ladder "
                   "(Pass above the US Army OPAT 1.60 m mark; Good / Excellent / Elite at Norwegian "
                   "conscripts' 25th / 50th / 75th, Aandstad 2023, CC BY; about the AFSPECWAR OFT minimum / "
                   "9 points / maximum) — " + PLAN61 + ", Part B §3.2")
BJ_SF_NOTE = ("Broad jump 1.8 / 2.0 / 2.2 / 2.4 -> 1.93 / 2.16 / 2.39 / 2.65 m: the AFSPECWAR OFT 8 / 9 / "
              "10-point marks (76 / 85 / 94 in, sex-neutral; a US government form, read from a third-party "
              "copy, min and max corroborated) and an Elite about the 80th percentile of AFSPECWAR candidates "
              "(Feeney 2023, facts) — " + PLAN61 + ", Part B §3.2")
PLANK_POLICE_NOTE = ("Plank 1:00 / 1:30 / 2:30 / 3:30 -> the operator plank 1:30 / 2:30 / 3:30 / 4:30: the old "
                     "ladder sat below every US entry-age minimum (USMC / Navy 1:10, AFT 1:20-1:30, PFRA "
                     "1:30-1:35); plan 61 Option A — " + PLAN61 + ", Part B §6.3")
PLANK_USMC_NOTE = ("Pass 1:03 -> 1:10: the USMC PFT plank minimum since 1 Jan 2022 (Marine Corps Times "
                   "2021-08-05 and two calculator sites; the order itself returned 403 to the app's research). "
                   "The app's unit source label now cites MCO 6100.13A (with changes), not 6100.13B, which no "
                   "source found; this site has no per-unit source label — " + PLAN61 + ", Part B §9.1-9.2")

OPS_MOVES = [
    # (pathway, benchmark, column, OLD, NEW, note or None)
    (PJ, '1.5-mile run', 'pass', '10:10', '10:20', IFT_NOTE.format(what='Pass 10:10 -> 10:20')),
    (PJ, PJ_SWIM_OLD, 'pass', '12:00', '12:30',
     IFT_NOTE.format(what="Pass 12:00 -> 12:30 (the PJ column; civilians enlist on the 9T5 column, 15:00)")),
    # The swim's id was pinned to `500_m_swim` by changeset 2 of
    # apply-military-top-tiers-2026-10-02.py, so this rename needs no id move.
    (PJ, PJ_SWIM_OLD, 'benchmark', PJ_SWIM_OLD, PJ_SWIM_NEW,
     f"Renamed '{PJ_SWIM_OLD}' -> '{PJ_SWIM_NEW}', the app's label (the IFT's PJ column); id kept "
     f"'500_m_swim' — {PLAN61}"),
    (PJ, PJ_PULL_OLD, 'benchmark', PJ_PULL_OLD, PJ_PULL_NEW,
     f"Renamed '{PJ_PULL_OLD}' -> '{PJ_PULL_NEW}', the app's label (the IFT allows 2 minutes; 1 minute is "
     f"for officers); id pinned to 'pull_ups_no_time' so stored results and shared inputs still match — "
     f"{PLAN61}"),
    (PJ, PJ_PULL_OLD, 'id', None, 'pull_ups_no_time', None),
    (PJ, 'Push-ups (2 min)', 'pass', 50, 40, IFT_NOTE.format(what='Pass 50 -> 40')),
    (PJ, 'Sit-ups (2 min)', 'pass', 54, 50, IFT_NOTE.format(what='Pass 54 -> 50')),
    # Planks.
    (USMC, 'Plank (front)', 'pass', '1:03', '1:10', PLANK_USMC_NOTE),
    *[(POLICE, 'Plank (front)', tier, o, n, PLANK_POLICE_NOTE)
      for tier, o, n in (('pass', '1:00', '1:30'), ('good', '1:30', '2:30'),
                         ('excellent', '2:30', '3:30'), ('elite', '3:30', '4:30'))],
    # Broad jumps.
    *[(unit, 'Broad Jump', tier, o, n, BJ_GENERAL_NOTE)
      for unit in (NAVY, POLICE)
      for tier, o, n in (('good', 1.9, 1.95), ('excellent', 2.1, 2.2), ('elite', 2.3, 2.4))],
    *[(unit, 'Broad Jump', tier, o, n, BJ_SF_NOTE)
      for unit in (SEAL, PJ, SFAS)
      for tier, o, n in (('pass', 1.8, 1.93), ('good', 2, 2.16), ('excellent', 2.2, 2.39), ('elite', 2.4, 2.65))],
]

# The app's cites, verbatim (tpf-app src/lib/operational_readiness.ts, read
# 2026-10-03 from its live ORS_PATHWAY_CONFIGS).
CITE_PLANK = ("TPF's operator plank (plan 61): Pass 1:30 = AFT 60 pts = PFRA women's minimum; Good ≈ AFT 78; "
              "Excellent ≈ AFT 97; Elite above every US maximum (3:24-3:45)")
CITE_HANG = ("TPF's own, checked 2026-10-03 (plan 61): no government table uses a dead hang; floors near "
             "untrained men's means (40-47 s), tops near elite swimmers' (≈ 95 s). CC BY, men only")
CITE_BJ_COMBAT = ("TPF's own, checked (plan 61): between the general and special-forces ladders; Elite = the "
                  "AFSPECWAR OFT maximum (2.39 m)")
CITE_BJ_GENERAL = ("Plan 61: Pass above the OPAT 1.60 m mark; Good / Excellent / Elite ≈ Norwegian conscripts' "
                   "25th / 50th / 75th (CC BY) ≈ the AFSPECWAR OFT minimum / 9 pts / maximum")
CITE_BJ_SF = ("AFSPECWAR OFT 8 / 9 / 10-point marks (76 / 85 / 94 in, sex-neutral); Elite 2.65 m ≈ the 80th "
              "percentile of AFSPECWAR candidates (plan 61)")
CITE_BJ_SWAT = ("TPF's own — checked 2026-10-03 (plan 61): Pass ≈ the AFSPECWAR OFT minimum (1.93 m); Elite ≈ "
                "the Swiss SOF operator mean + 0.6 SD (CC BY)")
CITE_PJ_SWIM = ("USAF IFT worksheet (10 Jan 2023), PJ column: 12:30 minimum. Civilians enlist on the 9T5 column, "
                "15:00")

CITES = [
    *[(u, 'Plank (front)', CITE_PLANK)
      for u in (PARA, RM, ABN, US_INF, SEAL, PJ, SFAS, UKSF, UK_INF, POLICE, SWAT)],
    *[(u, 'Dead hang (grip)', CITE_HANG)
      for u in (PARA, RM, USMC, ABN, NAVY, US_INF, SEAL, PJ, SFAS, UKSF, UK_INF)],
    (SWAT, 'Dead hang', CITE_HANG),  # SWAT's row is named "Dead hang" in both repositories
    *[(u, 'Broad Jump', CITE_BJ_COMBAT) for u in (PARA, RM, USMC, ABN)],
    *[(u, 'Broad Jump', CITE_BJ_GENERAL) for u in (NAVY, POLICE)],
    *[(u, 'Broad Jump', CITE_BJ_SF) for u in (SEAL, PJ, SFAS)],
    (SWAT, 'Broad Jump', CITE_BJ_SWAT),
    (PJ, PJ_SWIM_NEW, CITE_PJ_SWIM),
]
OPS_MOVES += [(u, b, 'cite', None, c, None) for u, b, c in CITES]

# The app's values, for `verify` (seconds / reps / metres).
PJ_APP = {
    '1.5-mile run': [620, 570, 540, 495],
    PJ_SWIM_NEW: [750, 660, 600, 570],
    PJ_PULL_NEW: [8, 12, 18, 25],
    'Push-ups (2 min)': [40, 70, 85, 105],
    'Sit-ups (2 min)': [50, 70, 90, 110],
}
BJ_APP = {
    **{u: [1.7, 1.95, 2.2, 2.4] for u in (NAVY, POLICE)},
    **{u: [1.93, 2.16, 2.39, 2.65] for u in (SEAL, PJ, SFAS)},
    **{u: [1.8, 2, 2.2, 2.4] for u in (PARA, RM, USMC, ABN)},
    SWAT: [1.9, 2.1, 2.3, 2.5],
}

OPS_README = (
    "2026-10-03, last (the missing standards researched — tpf-app plan 61 Part B; record: tpf-app "
    "docs/build/61_MISSING_STANDARDS_RESEARCH_2026-10-03.md section 8): USAF Pararescue on the IFT worksheet "
    "(10 Jan 2023, PJ column) — 1.5-mile Pass 10:20, 500 m swim Pass 12:30 (civilians enlist on the 15:00 "
    "column; the row renamed to the app's label), pull-ups '(2 min)' (renamed, id kept), push-ups Pass 40, "
    "sit-ups Pass 50. Broad jump one flat ladder per group: Navy PRT and police 1.70 / 1.95 / 2.20 / 2.40; "
    "SEAL, Pararescue and SFAS 1.93 / 2.16 / 2.39 / 2.65 (the AFSPECWAR OFT marks); combat units and SWAT "
    "kept. Police plank onto the operator plank 1:30 / 2:30 / 3:30 / 4:30; USMC plank Pass 1:03 -> 1:10. A "
    "new `cite` column holds the app's stated basis for every operator-plank, dead-hang and broad-jump row "
    "and the Pararescue swim (codegen does not read it; check:app-ors compares it). Script: "
    "scripts/apply-standards-research-2026-10-03.py."
)

# =============================================================== Lift ======

HYROX_M_OLD = ("Elite: Rappelt et al. 2026 (Front Physiol, CC BY), season-7 ELITE-division median (57:17) x 1.187; "
               "lower tiers TPF's own (provisional — need the owner's numbers). Numbers unchanged 2026-10-02")
HYROX_F_OLD = ("Elite: Rappelt et al. 2026 (Front Physiol, CC BY), season-7 ELITE-division median (1:03:22) x "
               "1.187 = 1:15:13, shown 1:15:00; lower tiers TPF's own (provisional), keeping the old spacing to "
               "elite. Rebuilt 2026-10-02")
HYROX_M_NEW = ("Elite: Rappelt et al. 2026 (Front Physiol, CC BY), season-7 ELITE-division median (57:17) x 1.187. "
               "Advanced: the median PRO racer of season 7 at Open weights (Rappelt Figure 2, x 0.93 for the "
               "lighter Open stations — TPF judgement) 1:14:28, kept 1:14:00. Intermediate: TPF's own, between "
               "the anchors. Experienced: Brandt et al. 2025 (Front Physiol, CC BY), recreational Open median "
               "86.5 min, kept 1:26:00. Novice / Beginner: 3.9 x TPF's own 5 km Novice / Beginner (25:00 / "
               "30:05; HYROX time scales with running, Rappelt and Brandt), rounded to 30 s: 1:37:30 / 1:57:30 "
               "(were 1:30:30 / 1:35:00). Mirrors tpf-app plan 61 Part A §3.4 (2026-10-03)")
HYROX_F_NEW = ("Women = men x 1.10 (Rappelt ELITE 1.106, PRO top-100 1.110), rounded to 30 s: Beginner 2:09:30, "
               "Novice 1:47:30, Experienced 1:34:30 (were 1:44:30 / 1:39:30 / 1:35:00); Intermediate 1:28:00 and "
               "Advanced 1:21:30 kept; Elite = Rappelt et al. 2026 (Front Physiol, CC BY) season-7 ELITE-division "
               "median (1:03:22) x 1.187 = 1:15:13, shown 1:15:00. Mirrors tpf-app plan 61 Part A §3.4 "
               "(2026-10-03)")
NO_STUDY = "TPF's own (expert-set); no openly licensed study exists to check it against (a gap)"
CINDY_M_NEW = ("TPF's own (expert-set), checked 2026-10-03 (tpf-app plan 61 Part A §4.5): Experienced 18 ≈ "
               "Kliszczewicz 2014's mixed group (17.8 rounds, CC BY); Advanced 23 ≈ Butcher 2015's Open / "
               "Regional athletes (23.3, CC BY-NC — cross-check only). Elite 25 kept (whether it should be 28 is "
               "an open owner question in the app's record)")
CINDY_F_NEW = ("TPF's own (expert-set), kept and checked 2026-10-03 (tpf-app plan 61 Part A §4.4): women's / men's "
               "0.83-0.91 (mean 0.88), inside the CrossFit evidence — Toledo 2021 0.94 and Mangine 2018 (Fight "
               "Gone Bad) 0.87, both CC BY; Kliszczewicz 2014's entry floor 0.71, CC BY; Forte 2022 0.95, CC "
               "BY-NC-ND, cross-check only. The levels agree: Forte's women averaged 15.7 rounds (Experienced 16); "
               "Kliszczewicz's women's floor is 10 (Beginner)")

BJ_SRC_OLD = "TPF's own (expert-set). No openly licensed population table has been checked yet (a gap)"
BJ_SRC_NEW = ("TPF's own (expert-set), checked 2026-10-03 (tpf-app plan 61 Part B §3.3) against Norwegian "
              "conscripts 2020 (Aandstad et al. 2023, Scand J Med Sci Sports, CC BY 4.0), US Army Ranger course "
              "candidates (Physiol Rep 2026, CC BY), Serbian police students (through a CC BY review, Healthcare "
              "2023), AFSPECWAR candidates (Feeney et al. 2023, facts) and China's National Student Physical "
              "Health Standard (2014 revision, a Ministry of Education standard). Men's Elite 285 -> 275 cm "
              "(2.85 m sat above every anchor); the women's ladder kept, about 0.80 x the men's at every tier")
BJ_REF_OLD = "TPF's own (expert-set); no permitted table checked yet (a gap)"
BJ_REF_NEW = ("TPF's own (expert-set), checked 2026-10-03 against permitted anchors (Norwegian conscripts, Aandstad "
              "2023, CC BY; US Army Ranger candidates, 2026, CC BY; China's student standard, 2014) — tpf-app "
              "plan 61 Part B §3.3")
BJ_NOTES_OLD = 'novice/intermediate/advanced interpolated from pass/good/excellent/elite (2026-08-07)'
BJ_NOTES_M_NEW = (BJ_NOTES_OLD + '; Elite 285 -> 275 cm 2026-10-03 (tpf-app plan 61 Part B §3.3: 2.85 m sat '
                  'above every anchor; the women\'s 220 keeps a 0.80 ratio)')

OHP_SRC_OLD = ("TPF's own ladder. No openly licensed population norm exists; checked only against published ratios "
               "(women's overhead press ≈ 0.5-0.6 of men's absolute — Nuzzo 2023 review) and the CrossFit sample "
               "of Meier, Rabel, Schmidt 2021 (Sports 9:80, CC BY)")
OHP_SRC_NEW = ("TPF's own ladder. No openly licensed population norm exists. Checked 2026-10-03 (tpf-app plan 61 "
               "Part C §2.2) against TPF's own bench / squat / deadlift ladders x the overhead-press ratios of the "
               "CrossFit samples of Meier, Rabel, Schmidt 2021 (Sports 9:80) and Dexheimer 2019 (both CC BY): the "
               "men's ladder reproduced to within 2.5 kg; the women's Beginner 20 -> 25 kg (the empty bar sat "
               "below every anchor, 23-30 kg). Women's ≈ 0.5-0.6 of men's absolute (Nuzzo 2023 review)")
OHP_LAUNCH_OLD = 'Owner-set; ratio checks only (2026-10-02)'
OHP_LAUNCH_NEW = 'Owner-set; ratio checks (2026-10-03)'
OHP_REF_OLD = "TPF's own; no openly licensed norm — ratio checks only (2026-10-02)"
OHP_REF_M_NEW = ("TPF's own; no openly licensed norm — checked 2026-10-03 against the CrossFit overhead-press "
                 "ratios (Meier 2021, Dexheimer 2019, CC BY): reproduced to within 2.5 kg, kept (tpf-app plan 61 "
                 "Part C §2.2)")
OHP_REF_F_NEW = ("TPF's own; no openly licensed norm — checked 2026-10-03 against the CrossFit overhead-press "
                 "ratios (Meier 2021, Dexheimer 2019, CC BY): Beginner 20 -> 25 kg, the empty bar sat below every "
                 "anchor (tpf-app plan 61 Part C §2.2). The HYROX, triathlete and bodybuilder pathways keep their "
                 "own 20 kg (an owner question in the app's record)")

BIKE_SRC_OLD = "TPF's own triathlete-pathway anchor (round 13: elite 27:00 M ≈ 44 km/h). No permitted anchor yet (a gap)"
BIKE_SRC_NEW = ("TPF's own triathlete-pathway anchor (round 13: elite 27:00 M ≈ 44 km/h), checked 2026-10-03 "
                "(tpf-app plan 61 Part C §4): Beginner at the US Army AFT 12 km stationary-bike pass time (2025 "
                "tables, a US government work) scaled to 20 km (Riegel exp 1.05) — men 45:10, women 49:32, so the "
                "women's Beginner moved 50:00 -> 49:30; Elite graded against the UCI Hour Records (facts) — men's "
                "Elite 74 % of world best; sex gap 1.11. Assumes a time-trial bike (or clip-on aero bars), a flat "
                "course and still air")
BIKE_LAUNCH_OLD = 'Owner-set (TPF)'
BIKE_LAUNCH_NEW = 'Owner-set (TPF); checked against permitted anchors (2026-10-03)'
BIKE_REF_OLD = "TPF's own (round 13); no permitted anchor yet (a gap)"
BIKE_REF_M_NEW = ("TPF's own (round 13); checked 2026-10-03: Beginner at the US Army AFT 12 km bike pass scaled to "
                  "20 km (45:10), Elite graded against the UCI Hour Records; a TT bike, flat, still air (tpf-app "
                  "plan 61 Part C §4)")
BIKE_REF_F_NEW = ("TPF's own (round 13); checked 2026-10-03: Beginner 50:00 -> 49:30, at the US Army AFT 12 km bike "
                  "pass scaled to 20 km (49:32 — it was 28 s easier than the Army minimum); a TT bike, flat, still "
                  "air (tpf-app plan 61 Part C §4)")

# (sheet, row key, column, OLD, NEW). The row key is the sheet's key columns:
# WOD_Standards (wod_id, sex); Standards (benchmark_id, sex);
# Benchmarks_Sourcing (benchmark_id,).
LIFT_MOVES = [
    ('WOD_Standards', ('hyrox_race', 'M'), 'pass', '1:35:00', '1:57:30'),
    ('WOD_Standards', ('hyrox_race', 'M'), 'novice', '1:30:30', '1:37:30'),
    ('WOD_Standards', ('hyrox_race', 'M'), 'data_source', HYROX_M_OLD, HYROX_M_NEW),
    ('WOD_Standards', ('hyrox_race', 'F'), 'pass', '1:44:30', '2:09:30'),
    ('WOD_Standards', ('hyrox_race', 'F'), 'novice', '1:39:30', '1:47:30'),
    ('WOD_Standards', ('hyrox_race', 'F'), 'good', '1:35:00', '1:34:30'),
    ('WOD_Standards', ('hyrox_race', 'F'), 'data_source', HYROX_F_OLD, HYROX_F_NEW),
    ('WOD_Standards', ('cindy', 'M'), 'data_source', NO_STUDY, CINDY_M_NEW),
    ('WOD_Standards', ('cindy', 'F'), 'data_source', NO_STUDY, CINDY_F_NEW),
    # TPF Benchmark's own six-tier broad jump (cm).
    ('Standards', ('broad_jump', 'M'), 'elite', 285, 275),
    ('Standards', ('broad_jump', 'M'), 'source_ref', BJ_REF_OLD, BJ_REF_NEW),
    ('Standards', ('broad_jump', 'M'), 'notes', BJ_NOTES_OLD, BJ_NOTES_M_NEW),
    ('Standards', ('broad_jump', 'F'), 'source_ref', BJ_REF_OLD, BJ_REF_NEW),
    ('Benchmarks_Sourcing', ('broad_jump',), 'data_source', BJ_SRC_OLD, BJ_SRC_NEW),
    ('Benchmarks_Sourcing', ('broad_jump',), 'launch_method', 'Expert-set (TPF)',
     'Expert-set (TPF); checked against permitted anchors (2026-10-03)'),
    # The HABS base table (beyond §8.6's list; the app is canonical).
    ('Standards', ('strict_press_1rm', 'F'), 'pass', 20, 25),
    ('Standards', ('strict_press_1rm', 'M'), 'source_ref', OHP_REF_OLD, OHP_REF_M_NEW),
    ('Standards', ('strict_press_1rm', 'F'), 'source_ref', OHP_REF_OLD, OHP_REF_F_NEW),
    ('Benchmarks_Sourcing', ('strict_press_1rm',), 'data_source', OHP_SRC_OLD, OHP_SRC_NEW),
    ('Benchmarks_Sourcing', ('strict_press_1rm',), 'launch_method', OHP_LAUNCH_OLD, OHP_LAUNCH_NEW),
    ('Standards', ('bike_20k', 'F'), 'pass', '50:00', '49:30'),
    ('Standards', ('bike_20k', 'M'), 'source_ref', BIKE_REF_OLD, BIKE_REF_M_NEW),
    ('Standards', ('bike_20k', 'F'), 'source_ref', BIKE_REF_OLD, BIKE_REF_F_NEW),
    ('Benchmarks_Sourcing', ('bike_20k',), 'data_source', BIKE_SRC_OLD, BIKE_SRC_NEW),
    ('Benchmarks_Sourcing', ('bike_20k',), 'launch_method', BIKE_LAUNCH_OLD, BIKE_LAUNCH_NEW),
    # The 40 km is derived from the 20 km (Riegel 1.05, 10 s steps), as in the
    # app: 49:30 x 2.070 = 1:42:29 -> 1:42:30 (sheet units mm:ss, so 102:30).
    ('Standards', ('bike_40k', 'F'), 'pass', '103:30', '102:30'),
]

LIFT_README_KEY = 'Change 2026-10-03'
LIFT_README = (
    "Mirrors the TPF app's plan 61 (tpf-app docs/build/61_MISSING_STANDARDS_RESEARCH_2026-10-03.md section 8). "
    "HYROX race: men Beginner 1:57:30 and Novice 1:37:30 (3.9 x TPF's 5 km), women Beginner 2:09:30, Novice "
    "1:47:30, Experienced 1:34:30 (men x 1.10); every tier now has a permitted basis. Cindy kept, now checked. "
    "TPF Benchmark's own broad jump: men's Elite 285 -> 275 cm, relabelled TPF's own, now checked. The HABS base "
    "table follows the app: women's overhead press Beginner 20 -> 25 kg, women's 20 km TT Beginner 50:00 -> "
    "49:30 (40 km derived: 1:42:30). Script: scripts/apply-standards-research-2026-10-03.py."
)

# ============================================================== helpers ====


def col_for(idx, name):
    """The column whose header is `name` or starts with `name (` (the tier
    columns carry their score, e.g. 'pass (50%)')."""
    for k, v in idx.items():
        if k == name or k.startswith(name + ' ('):
            return v
    raise KeyError(name)


def check_drift(label, current, old, new):
    if current == old or current == new:
        return None
    return f'{label}: found {current!r}, expected {old!r} (or {new!r} on a re-run)'


# ---------------------------------------------------------- Operator ------

ops_wb = openpyxl.load_workbook(OPS)
BOOK_NAME = {id(ops_wb): OPS.name}
ws = ops_wb['Standards']
HEADER = 2
idx = header_index(ws, HEADER)
if 'cite' not in idx:
    put(ws, HEADER, ws.max_column + 1, 'cite')
    idx = header_index(ws, HEADER)

first_name = {}  # (pathway, any name the row has had) -> (pathway, its first name)


def row_key(pathway, bench):
    return first_name.get((pathway, bench), (pathway, bench))


for pathway, bench, col, old, new, _ in OPS_MOVES:
    if col == 'benchmark':
        assert row_key(pathway, old) == row_key(pathway, bench), f'{(pathway, bench)}: rename expects {old!r}'
        first_name[(pathway, new)] = row_key(pathway, bench)

ops_rows = {}
for r in range(HEADER + 1, ws.max_row + 1):
    pathway, bench = ws.cell(row=r, column=idx['pathway']).value, ws.cell(row=r, column=idx['benchmark']).value
    if pathway is None:
        continue
    key = row_key(pathway, bench)
    assert key not in ops_rows, f'duplicate row {key} (sheet name {bench!r})'
    ops_rows[key] = r

seen = set()
drift = []
for pathway, bench, col, old, new, _ in OPS_MOVES:
    key = row_key(pathway, bench)
    assert (key, col) not in seen, f'{key} / {col} is moved twice in this changeset'
    seen.add((key, col))
    if key not in ops_rows:
        drift.append(f'{key}: no such row on the Operator sheet')
        continue
    d = check_drift(f'Operator {key} / {col}', ws.cell(row=ops_rows[key], column=idx[col]).value, old, new)
    if d:
        drift.append(d)

# ------------------------------------------------------------- Lift -------

hrs_wb = openpyxl.load_workbook(HRS)
BOOK_NAME[id(hrs_wb)] = HRS.name
SHEET_KEYS = {'WOD_Standards': ('wod_id', 'sex'), 'Standards': ('benchmark_id', 'sex'),
              'Benchmarks_Sourcing': ('benchmark_id',)}
lift_rows = {}  # sheet -> {key tuple: row}
lift_idx = {}
for sheet, keys in SHEET_KEYS.items():
    s = hrs_wb[sheet]
    i = header_index(s, 2)
    lift_idx[sheet] = i
    lift_rows[sheet] = {}
    for r in range(3, s.max_row + 1):
        k = tuple(s.cell(row=r, column=i[c]).value for c in keys)
        if k[0] is None:
            continue
        assert k not in lift_rows[sheet], f'{sheet}: duplicate row {k}'
        lift_rows[sheet][k] = r

for sheet, key, col, old, new in LIFT_MOVES:
    s = hrs_wb[sheet]
    if key not in lift_rows[sheet]:
        drift.append(f'{sheet} {key}: no such row')
        continue
    d = check_drift(f'{sheet} {key} / {col}', s.cell(row=lift_rows[sheet][key], column=col_for(lift_idx[sheet], col)).value,
                    old, new)
    if d:
        drift.append(d)

if drift:
    raise SystemExit('A workbook has drifted from this record; nothing was saved:\n  ' + '\n  '.join(drift))

# ------------------------------------------------------------ apply -------

notes = {}
for pathway, bench, col, old, new, note in OPS_MOVES:
    key = row_key(pathway, bench)
    put(ws, ops_rows[key], idx[col], new)
    if note:
        notes.setdefault(key, [])
        if note not in notes[key]:
            notes[key].append(note)
for key, ns in notes.items():
    r = ops_rows[key]
    source = ws.cell(row=r, column=idx['source']).value or ''
    for n in ns:
        if n not in source:
            source = f'{source}; {n}' if source else n
    put(ws, r, idx['source'], source)

for sheet, key, col, old, new in LIFT_MOVES:
    put(hrs_wb[sheet], lift_rows[sheet][key], col_for(lift_idx[sheet], col), new)

# ----------------------------------------------------------- verify -------

errors = []


def ops_ladder(pathway, bench):
    r = ops_rows[row_key(pathway, bench)]
    return [ws.cell(row=r, column=idx[c]).value for c in ('pass', 'good', 'excellent', 'elite')]


for bench, want in PJ_APP.items():
    got = [to_sec(v) for v in ops_ladder(PJ, bench)]
    if got != want:
        errors.append(f'Operator {PJ} / {bench}: the sheet reads {got}, the app {want}')
for unit, want in BJ_APP.items():
    got = ops_ladder(unit, 'Broad Jump')
    if got != want:
        errors.append(f'Operator {unit} / Broad Jump: the sheet reads {got}, the app {want}')
for unit in (POLICE, USMC):
    got = [to_sec(v) for v in ops_ladder(unit, 'Plank (front)')]
    want = [90, 150, 210, 270] if unit == POLICE else [70, 180, 210, 270]
    if got != want:
        errors.append(f'Operator {unit} / Plank (front): the sheet reads {got} s, the app {want} s')
# Every moved Operator row is strictly ordered in its direction.
for key in {row_key(p, b) for p, b, *_ in OPS_MOVES}:
    r = ops_rows[key]
    t = [to_sec(ws.cell(row=r, column=idx[c]).value) for c in ('pass', 'good', 'excellent', 'elite')]
    lower = 'lower' in str(ws.cell(row=r, column=idx['direction']).value)
    if not all((a > b) if lower else (a < b) for a, b in zip(t, t[1:])):
        errors.append(f'Operator {key}: not strictly ordered: {t}')
# Every mirrored cite is on a row the sheet holds, once.
cite_rows = [row_key(u, b) for u, b, _ in CITES]
if len(set(cite_rows)) != len(cite_rows):
    errors.append('a cite row is listed twice')

wod = hrs_wb['WOD_Standards']
wi = lift_idx['WOD_Standards']
TIERS6 = ('pass', 'novice', 'good', 'intermediate', 'advanced', 'elite')


def wod_ladder(key):
    r = lift_rows['WOD_Standards'][key]
    return [to_sec(wod.cell(row=r, column=col_for(wi, t)).value) for t in TIERS6]


std = hrs_wb['Standards']
si = lift_idx['Standards']


def std_cell(key, tier):
    return std.cell(row=lift_rows['Standards'][key], column=col_for(si, tier)).value


# HYROX: Beginner / Novice = 3.9 x the sheet's own 5 km Beginner / Novice, to
# the nearest 30 s; women = men x 1.10 to the nearest 30 s at every tier.
hm, hf = wod_ladder(('hyrox_race', 'M')), wod_ladder(('hyrox_race', 'F'))
for i, tier in ((0, 'pass'), (1, 'novice')):
    want = round_half_up(3.9 * to_sec(std_cell(('run_5k', 'M'), tier)), 30)
    if hm[i] != want:
        errors.append(f'HYROX men {tier}: {hm[i]} s, 3.9 x the 5 km {tier} gives {want} s')
for i, tier in enumerate(TIERS6):
    want = round_half_up(hm[i] * 1.10, 30)
    if hf[i] != want:
        errors.append(f'HYROX women {tier}: {hf[i]} s, men x 1.10 gives {want} s')
# The six-tier broad jump: women about 0.80 x men at every tier.
bm = [std_cell(('broad_jump', 'M'), t) for t in TIERS6]
bf = [std_cell(('broad_jump', 'F'), t) for t in TIERS6]
for t, m, f in zip(TIERS6, bm, bf):
    if not 0.78 <= f / m <= 0.82:
        errors.append(f'broad jump {t}: women {f} / men {m} = {f / m:.3f}, outside 0.78-0.82')
# The vestigial four-tier `excellent` (never scored or shown on a six-tier row)
# keeps its place among the six tiers: no move here may pass it.
TIER_COLS = set(TIERS6)
moved_tiers = {}
for sheet, key, col, old, new in LIFT_MOVES:
    if col in TIER_COLS:
        moved_tiers.setdefault((sheet, key), {})[col] = old
for (sheet, key), olds in moved_tiers.items():
    s, i = hrs_wb[sheet], lift_idx[sheet]
    r = lift_rows[sheet][key]
    now = {t: to_sec(s.cell(row=r, column=col_for(i, t)).value) for t in TIERS6}
    before = {**now, **{t: to_sec(v) for t, v in olds.items()}}
    ex = to_sec(s.cell(row=r, column=col_for(i, 'excellent')).value)
    if sum(v < ex for v in before.values()) != sum(v < ex for v in now.values()):
        errors.append(f'{sheet} {key}: a move passed the vestigial excellent ({ex})')

if errors:
    raise SystemExit('A check of the finished sheets failed; nothing was saved:\n  ' + '\n  '.join(errors))

# ------------------------------------------------------------ README -------

rd = ops_wb['README']
if not any(c.value == OPS_README for c in rd['A']):
    put(rd, rd.max_row + 2, 1, OPS_README)

rm = hrs_wb['Read me']
if not any(c.value == LIFT_README for c in rm['B']):
    r = rm.max_row + 1
    put(rm, r, 1, LIFT_README_KEY)
    put(rm, r, 2, LIFT_README)

# ------------------------------------------------------------- save -------
# Only a workbook something changed in is saved, so a re-run is a true no-op.

for wb, path in ((ops_wb, OPS), (hrs_wb, HRS)):
    if path.name in dirty:
        wb.save(path)

print(f'{len(changes)} cell change(s); saved: {", ".join(sorted(dirty)) or "nothing"}')
for c in changes:
    print('  ' + c)
